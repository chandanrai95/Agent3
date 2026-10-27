import asyncio
import base64
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import pymupdf as fitz
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook

from app.config.setting import VISION_MODEL, MAX_IMAGES_PER_DOCUMENT

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"} | IMAGE_EXTENSIONS

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150
TABLE_CHUNK_SIZE = 1500
# a page with less text than this is treated as scanned and read with the vision model
SCANNED_PAGE_MIN_CHARS = 50
# skip icons, bullets and logos
MIN_IMAGE_SIDE = 120
VISION_CONCURRENCY = 4
NO_CONTENT = "NO_CONTENT"

IMAGE_PROMPT = f"""You are indexing an image from a document for a search engine.
- Transcribe all visible text verbatim.
- If it contains a table, render it as a GitHub markdown table.
- For charts, state the chart type, axes, and the key values and trends.
- For diagrams or photos, describe the components and what they show.
Be factual and concise. If the image is purely decorative (logo, icon, divider) reply exactly {NO_CONTENT}."""

SCANNED_PAGE_PROMPT = f"""This is a scanned document page. Transcribe its content for a search index.
- Transcribe all text verbatim in reading order, keep headings.
- Render every table as a GitHub markdown table.
- Describe charts, photos, stamps or signatures briefly as [Figure: ...].
Output only the page content. If the page is blank reply exactly {NO_CONTENT}."""

MIME_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}


@dataclass
class ImageJob:
    data: bytes
    mime: str
    content_type: str  # "image" or "scanned_page"
    metadata: dict = field(default_factory=dict)


@dataclass
class ParsedDocument:
    # markdown segments, tables are inline markdown tables
    segments: list[Document] = field(default_factory=list)
    images: list[ImageJob] = field(default_factory=list)
    skipped_images: int = 0

    def add_image(self, job: ImageJob):
        if len(self.images) >= MAX_IMAGES_PER_DOCUMENT:
            self.skipped_images += 1
            return
        self.images.append(job)


# -----------------------------
# MARKDOWN HELPERS
# -----------------------------
def _cell(value) -> str:
    return "" if value is None else str(value).replace("|", "\\|").replace("\n", " ").strip()


def rows_to_markdown(rows: list[list]) -> str:
    rows = [r for r in rows if any(_cell(c) for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# -----------------------------
# PARSERS (sync, run in a thread)
# -----------------------------
def _pixmap_png(doc: fitz.Document, xref: int) -> Optional[bytes]:
    pix = fitz.Pixmap(doc, xref)
    if pix.width < MIN_IMAGE_SIDE or pix.height < MIN_IMAGE_SIDE:
        return None
    if pix.n - pix.alpha >= 4:  # CMYK -> RGB
        pix = fitz.Pixmap(fitz.csRGB, pix)
    return pix.tobytes("png")


def _text_blocks(page: fitz.Page, clip=None, hidden: bool = False) -> list[tuple[fitz.Rect, str]]:
    """Visible text blocks (or only the invisible OCR layer when hidden=True),
    without fake-bold duplicates (same span drawn twice with a small offset)."""
    blocks, kept = [], []
    for block in page.get_text("dict", clip=clip, sort=True)["blocks"]:
        if block["type"] != 0:
            continue
        lines = []
        for line in block["lines"]:
            spans = []
            for span in line["spans"]:
                # alpha 0 = invisible text, usually an OCR layer drawn over the page
                if (span.get("alpha", 255) == 0) != hidden:
                    continue
                x, y = span["bbox"][:2]
                if any(t == span["text"] and abs(x - kx) < 3 and abs(y - ky) < 3 for t, kx, ky in kept):
                    continue
                kept.append((span["text"], x, y))
                spans.append(span["text"])
            if "".join(spans).strip():
                lines.append("".join(spans).strip())
        if lines:
            blocks.append((fitz.Rect(block["bbox"]), "\n".join(lines)))
    return blocks


def _extract_table(page: fitz.Page, table, hidden: bool = False) -> list[list[str]]:
    # table.extract() duplicates glyphs of fake-bold text ("TTaabbllee"), read cells as deduped text
    return [
        [" ".join(t for _, t in _text_blocks(page, clip=cell, hidden=hidden)) if cell else "" for cell in row.cells]
        for row in table.rows
    ]


def _char_count(blocks: list[tuple[fitz.Rect, str]]) -> int:
    return sum(len(t.strip()) for _, t in blocks)


def parse_pdf(file_path: Path) -> ParsedDocument:
    parsed = ParsedDocument()
    seen_xrefs = set()

    with fitz.open(file_path) as doc:
        for page in doc:
            page_no = page.number + 1
            meta = {"page": page_no}

            blocks = _text_blocks(page)
            hidden = False
            if _char_count(blocks) < SCANNED_PAGE_MIN_CHARS:
                # no visible text: use an existing OCR layer if there is one, else read the page with vision
                ocr_blocks = _text_blocks(page, hidden=True)
                if _char_count(ocr_blocks) >= SCANNED_PAGE_MIN_CHARS:
                    blocks, hidden = ocr_blocks, True
                else:
                    if page.get_images() or page.get_drawings():
                        pix = page.get_pixmap(dpi=150)
                        parsed.add_image(ImageJob(pix.tobytes("png"), "image/png", "scanned_page", meta))
                    continue

            # the default table finder needs ruling lines, skip the slow search on pages without any
            tables = page.find_tables().tables if page.get_drawings() else []
            table_rows = [_extract_table(page, t, hidden) for t in tables]
            # one column "tables" are boxed text or a table of contents, keep them as text
            tables = [t for t, rows in zip(tables, table_rows) if max((len(r) for r in rows), default=0) > 1]
            table_rows = [rows for rows in table_rows if max((len(r) for r in rows), default=0) > 1]
            table_rects = [fitz.Rect(t.bbox) for t in tables]

            # text blocks outside tables so table cells aren't indexed twice as broken prose
            text = "\n".join(
                t for rect, t in blocks
                if not any(rect.intersects(r) for r in table_rects)
            )
            if text:
                parsed.segments.append(Document(page_content=text, metadata={**meta, "content_type": "text"}))

            for rows in table_rows:
                markdown = rows_to_markdown(rows)
                if markdown:
                    parsed.segments.append(Document(page_content=markdown, metadata={**meta, "content_type": "table"}))

            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen_xrefs:  # logos repeated on every page
                    continue
                seen_xrefs.add(xref)
                try:
                    png = _pixmap_png(doc, xref)
                except Exception:
                    continue
                if png:
                    parsed.add_image(ImageJob(png, "image/png", "image", meta))

    return parsed


def parse_docx(file_path: Path) -> ParsedDocument:
    parsed = ParsedDocument()
    doc = DocxDocument(str(file_path))

    # walk the body in order so tables stay next to the text that introduces them
    parts = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = DocxParagraph(child, doc)
            if p.text.strip():
                style = (p.style.name or "") if p.style is not None else ""
                level = re.match(r"Heading (\d)", style)
                parts.append(("#" * int(level.group(1)) + " " if level else "") + p.text.strip())
        elif tag == "tbl":
            table = DocxTable(child, doc)
            parts.append(rows_to_markdown([[c.text for c in row.cells] for row in table.rows]))
    if parts:
        parsed.segments.append(Document(page_content="\n\n".join(parts), metadata={"content_type": "text"}))

    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        if part.content_type in MIME_TYPES.values():
            parsed.add_image(ImageJob(part.blob, part.content_type, "image"))

    return parsed


def parse_xlsx(file_path: Path) -> ParsedDocument:
    parsed = ParsedDocument()
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        markdown = rows_to_markdown([list(r) for r in sheet.iter_rows(values_only=True)])
        if markdown:
            parsed.segments.append(Document(page_content=markdown, metadata={"sheet": sheet.title, "content_type": "table"}))
    workbook.close()
    return parsed


def _read_text(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return file_path.read_text(errors="ignore")


def parse_csv(file_path: Path) -> ParsedDocument:
    rows = list(csv.reader(_read_text(file_path).splitlines()))
    markdown = rows_to_markdown(rows)
    segments = [Document(page_content=markdown, metadata={"content_type": "table"})] if markdown else []
    return ParsedDocument(segments=segments)


def parse_text(file_path: Path) -> ParsedDocument:
    text = _read_text(file_path).strip()
    return ParsedDocument(segments=[Document(page_content=text, metadata={"content_type": "text"})] if text else [])


def parse_image(file_path: Path) -> ParsedDocument:
    parsed = ParsedDocument()
    parsed.add_image(ImageJob(file_path.read_bytes(), MIME_TYPES[file_path.suffix.lower()], "image"))
    return parsed


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".csv": parse_csv,
    ".txt": parse_text,
    ".md": parse_text,
    **{ext: parse_image for ext in IMAGE_EXTENSIONS},
}


def parse_file(file_path: Path) -> ParsedDocument:
    ext = file_path.suffix.lower()
    if ext not in PARSERS:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    return PARSERS[ext](file_path)


# -----------------------------
# VISION (images, scanned pages)
# -----------------------------
async def describe_images(jobs: list[ImageJob]) -> tuple[list[Document], int]:
    """Turns images into searchable markdown. Returns (segments, failed_count)."""
    if not jobs:
        return [], 0

    llm = ChatOpenAI(model=VISION_MODEL, temperature=0)
    semaphore = asyncio.Semaphore(VISION_CONCURRENCY)

    async def describe(job: ImageJob) -> Optional[Document]:
        prompt = SCANNED_PAGE_PROMPT if job.content_type == "scanned_page" else IMAGE_PROMPT
        image_url = f"data:{job.mime};base64,{base64.b64encode(job.data).decode()}"
        async with semaphore:
            response = await llm.ainvoke([HumanMessage(content=[
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": image_url}},
            ])])
        content = (response.content or "").strip()
        if not content or content == NO_CONTENT:
            return None
        return Document(page_content=content, metadata={**job.metadata, "content_type": job.content_type})

    results = await asyncio.gather(*(describe(j) for j in jobs), return_exceptions=True)
    failed = [r for r in results if isinstance(r, Exception)]
    for e in failed:
        print(f"WARNING: image description failed: {e}")
    return [r for r in results if isinstance(r, Document)], len(failed)


# -----------------------------
# CHUNKING
# -----------------------------
_text_splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
_TABLE_LINE = re.compile(r"^\s*\|.*\|\s*$")


def _split_blocks(text: str) -> list[tuple[str, str]]:
    """Splits markdown into ("table" | "text", content) blocks."""
    blocks, current, current_kind = [], [], None
    for line in text.splitlines():
        kind = "table" if _TABLE_LINE.match(line) else "text"
        if kind != current_kind and current:
            blocks.append((current_kind, "\n".join(current)))
            current = []
        current_kind = kind
        current.append(line)
    if current:
        blocks.append((current_kind, "\n".join(current)))
    return [(k, c) for k, c in blocks if c.strip()]


def split_table(markdown: str) -> list[str]:
    """Row-group chunks that each repeat the header, so every chunk stays a readable table."""
    lines = markdown.splitlines()
    if len(lines) <= 2:
        return [markdown]
    header, rows = lines[:2], lines[2:]
    chunks, current = [], []
    for row in rows:
        if current and len("\n".join(header + current + [row])) > TABLE_CHUNK_SIZE:
            chunks.append("\n".join(header + current))
            current = []
        current.append(row)
    chunks.append("\n".join(header + current))
    return chunks


def _label(source_name: str, metadata: dict) -> str:
    where = f", page {metadata['page']}" if metadata.get("page") else ""
    where += f", sheet {metadata['sheet']}" if metadata.get("sheet") else ""
    kind = {"table": "Table", "image": "Image", "scanned_page": "Scanned page"}.get(metadata.get("content_type"))
    return f"[{kind} from {source_name}{where}]\n" if kind else ""


def chunk_segments(segments: list[Document], source_name: str) -> list[Document]:
    chunks = []
    for segment in segments:
        segment_type = segment.metadata.get("content_type", "text")
        for kind, block in _split_blocks(segment.page_content):
            metadata = {**segment.metadata, "content_type": "table" if kind == "table" else segment_type}
            # a label gives tables and image descriptions context of their own for embedding
            label = _label(source_name, metadata)
            pieces = split_table(block) if kind == "table" else _text_splitter.split_text(block)
            chunks += [Document(page_content=label + p, metadata=metadata) for p in pieces if p.strip()]
    return chunks
